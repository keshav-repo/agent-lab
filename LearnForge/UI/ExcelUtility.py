import sys
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, ttk

from KnowledgeService import update_metadata_using_excel

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment

from fs_utils import UploadFiles
from sqlLiteDB import get_LearningEntity_join_alias

COLUMNS = (
    "id",
    "text",
    "category",
    "subcategory",
    "topic",
    "subtopic",
    "concept",
    "tags",
    "aliases",
)


def _aliases_cell(aliases):
    if not aliases:
        return ""
    return "\n".join(f"• {alias}" for alias in aliases)


def _row_values(entity):
    return (
        entity.id,
        entity.text,
        entity.category,
        entity.subcategory or "",
        entity.topic or "",
        entity.subtopic or "",
        entity.concept or "",
        ", ".join(entity.tags),
        _aliases_cell(entity.aliases),
    )


class ExcelUtility(ttk.Frame):
    def __init__(self, master, on_back=None):
        super().__init__(master, padding=24)
        self.on_back = on_back
        self._build()

    def _build(self):
        header = ttk.Frame(self)
        header.pack(fill=tk.X, pady=(0, 24))
        if self.on_back:
            ttk.Button(
                header,
                text="← Back",
                style="Back.TButton",
                command=self.on_back,
            ).pack(side=tk.LEFT)

        ttk.Label(self, text="Excel Utility", style="Title.TLabel").pack(pady=(12, 6))
        ttk.Label(
            self,
            text="Choose an Excel action.",
            style="Subtitle.TLabel",
        ).pack(pady=(0, 28))

        actions = ttk.Frame(self)
        actions.pack()

        ttk.Button(
            actions,
            text="Download Learning Entity",
            command=self.download_learning_entity,
        ).pack(fill=tk.X, pady=8, ipadx=48)

        ttk.Button(
            actions,
            text="Update Metadata",
            style="Secondary.TButton",
            command=self.update_metadata,
        ).pack(fill=tk.X, pady=8, ipadx=48)

        self.status_label = ttk.Label(self, text="", style="Muted.TLabel")
        self.status_label.pack(pady=(28, 0))

    def download_learning_entity(self):
        path = filedialog.asksaveasfilename(
            parent=self.winfo_toplevel(),
            title="Save Learning Entities",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="learning_entities.xlsx",
        )
        if not path:
            self.status_label.config(text="Download cancelled")
            return

        entities = get_LearningEntity_join_alias()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Learning Entities"
        sheet.append(list(COLUMNS))
        for entity in entities:
            sheet.append(list(_row_values(entity)))
        aliases_col = COLUMNS.index("aliases") + 1
        sheet.column_dimensions[sheet.cell(1, aliases_col).column_letter].width = 40
        for row in sheet.iter_rows(min_row=2, min_col=aliases_col, max_col=aliases_col):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        workbook.save(path)
        self.status_label.config(text=f"Downloaded {len(entities)} entities")

    def update_metadata(self):
        file_path = filedialog.askopenfilename(
            parent=self.winfo_toplevel(),
            title="Select Excel file",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        )
        if not file_path:
            self.status_label.config(text="No file selected.")
            return

        UploadFiles([file_path])
        # call meta data upload
        update_metadata_using_excel()
        self.status_label.config(text="Excel file uploaded successfully.")
