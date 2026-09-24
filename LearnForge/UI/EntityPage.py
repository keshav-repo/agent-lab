import sys
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment
from pdfUtil import convert_to_pdf

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from KnowledgeService import get_LearningItems, get_pdf_content
from theme import apply_theme, colors, toggle_label, toggle_theme

COLUMNS = (
    "id",
    "text",
    "category",
    "subcategory",
    "topic",
    "subtopic",
    "concept",
    "tags",
    "aliasCount",
)
FILTER_FIELDS = ("category", "subcategory", "topic", "subtopic")
ALL = "All"


def _entity_values(entity):
    return (
        entity.id,
        entity.text,
        entity.category,
        entity.subcategory or "",
        entity.topic or "",
        entity.subtopic or "",
        entity.concept or "",
        ", ".join(entity.tags),
        entity.aliasCount or 0,
    )


def _search_blob(entity):
    return " ".join(str(value) for value in _entity_values(entity)).lower()


class EntityPage(ttk.Frame):
    def __init__(self, master, on_back=None):
        super().__init__(master, padding=8)
        self.on_back = on_back
        self.entities = []
        self.filter_vars = {}
        self.filter_boxes = {}
        self._updating_filters = False
        self._build()
        self.refresh()

    def _build(self):
        frame = self

        header = ttk.Frame(frame)
        header.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 12))
        if self.on_back:
            ttk.Button(
                header,
                text="← Back",
                style="Back.TButton",
                command=self.on_back,
            ).pack(side=tk.LEFT)
        self.count_label = ttk.Label(header, text="", style="Muted.TLabel")
        self.count_label.pack(side=tk.RIGHT, padx=(0, 130))

        toolbar = ttk.Frame(frame)
        toolbar.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(0, 8))

        ttk.Label(toolbar, text="Search", style="Muted.TLabel").pack(side=tk.LEFT, padx=(0, 8))
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(toolbar, textvariable=self.search_var)
        search_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.search_var.trace_add("write", lambda *_args: self.apply_filter())

        ttk.Button(toolbar, text="Clear", style="Secondary.TButton", command=self.clear_filters).pack(
            side=tk.LEFT, padx=(8, 0)
        )
        ttk.Button(toolbar, text="Refresh", style="Secondary.TButton", command=self.refresh).pack(
            side=tk.LEFT, padx=(8, 0)
        )
        ttk.Button(toolbar, text="Download PDF", style="Secondary.TButton", command=self.download_pdf).pack(
            side=tk.LEFT, padx=(8, 0)
        )
        ttk.Button(toolbar, text="Download Excel", style="Secondary.TButton", command=self.download_excel).pack(
            side=tk.LEFT, padx=(8, 0)
        )

        filters = ttk.Frame(frame)
        filters.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(0, 8))
        for field in FILTER_FIELDS:
            ttk.Label(filters, text=field.title(), style="Muted.TLabel").pack(
                side=tk.LEFT, padx=(0, 4)
            )
            var = tk.StringVar(value=ALL)
            box = ttk.Combobox(filters, textvariable=var, state="readonly", width=20)
            box.pack(side=tk.LEFT, padx=(0, 12))
            var.trace_add("write", lambda *_args, current=field: self._on_filter_change(current))
            self.filter_vars[field] = var
            self.filter_boxes[field] = box

        self.tree = ttk.Treeview(frame, columns=COLUMNS, show="headings")
        self.refresh_theme()
        for col in COLUMNS:
            self.tree.heading(col, text=col.replace("_", " ").title())
            self.tree.column(col, width=120, stretch=True)
        self.tree.column("id", width=50, stretch=False)
        self.tree.column("text", width=360)
        self.tree.column("aliasCount", width=90, stretch=False)

        yscroll = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=self.tree.yview)
        xscroll = ttk.Scrollbar(frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)

        self.tree.grid(row=3, column=0, sticky="nsew")
        yscroll.grid(row=3, column=1, sticky="ns")
        xscroll.grid(row=4, column=0, sticky="ew")
        frame.rowconfigure(3, weight=1)
        frame.columnconfigure(0, weight=1)

    def clear_filters(self):
        self._updating_filters = True
        self.search_var.set("")
        for var in self.filter_vars.values():
            var.set(ALL)
        self._updating_filters = False
        self._sync_filter_options()
        self.apply_filter()

    def refresh(self):
        self.entities = get_LearningItems()
        self._sync_filter_options()
        self.apply_filter()

    def _visible_ids(self):
        return [
            int(self.tree.item(item_id, "values")[0])
            for item_id in self.tree.get_children()
            if self.tree.item(item_id, "values")
        ]

    def download_pdf(self):
        path = filedialog.asksaveasfilename(
            parent=self.winfo_toplevel(),
            title="Save PDF",
            defaultextension=".pdf",
            filetypes=[("PDF files", "*.pdf")],
            initialfile="learning_entities.pdf",
        )
        if not path:
            return
        content = get_pdf_content(self._visible_ids())
        convert_to_pdf(content, path)

    def download_excel(self):
        path = filedialog.asksaveasfilename(
            parent=self.winfo_toplevel(),
            title="Save Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="learning_entities.xlsx",
        )
        if not path:
            return
        entities = get_pdf_content(self._visible_ids())
        columns = (
            "id",
            "text",
            "category",
            "subcategory",
            "topic",
            "subtopic",
            "concept",
            "tags",
            "aliases",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Learning Entities"
        sheet.append(list(columns))
        for entity in entities:
            aliases = "\n".join(f"• {alias}" for alias in entity.aliases if alias)
            sheet.append(
                [
                    entity.id,
                    entity.text,
                    entity.category,
                    entity.subcategory or "",
                    entity.topic or "",
                    entity.subtopic or "",
                    entity.concept or "",
                    ", ".join(entity.tags),
                    aliases,
                ]
            )
        aliases_col = columns.index("aliases") + 1
        sheet.column_dimensions[sheet.cell(1, aliases_col).column_letter].width = 40
        for row in sheet.iter_rows(min_row=2, min_col=aliases_col, max_col=aliases_col):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        workbook.save(path)

    def refresh_theme(self):
        palette = colors()
        self.tree.tag_configure(
            "odd",
            background=palette["row_odd"],
            foreground=palette["text"],
        )
        self.tree.tag_configure(
            "even",
            background=palette["row_even"],
            foreground=palette["text"],
        )
        self.apply_filter()

    def _selected_filters(self, until=None):
        selected = {}
        for field in FILTER_FIELDS:
            if field == until:
                break
            value = self.filter_vars[field].get()
            if value and value != ALL:
                selected[field] = value
        return selected

    def _matching_entities(self, until=None):
        selected = self._selected_filters(until=until)
        return [
            entity
            for entity in self.entities
            if all(getattr(entity, field) == value for field, value in selected.items())
        ]

    def _sync_filter_options(self):
        self._updating_filters = True
        for field in FILTER_FIELDS:
            values = sorted(
                {
                    getattr(entity, field)
                    for entity in self._matching_entities(until=field)
                    if getattr(entity, field)
                }
            )
            self.filter_boxes[field]["values"] = (ALL, *values)
            if self.filter_vars[field].get() not in (ALL, *values):
                self.filter_vars[field].set(ALL)
        self._updating_filters = False

    def _on_filter_change(self, _field):
        if self._updating_filters:
            return
        self._sync_filter_options()
        self.apply_filter()

    def apply_filter(self):
        query = self.search_var.get().strip().lower()
        selected = self._selected_filters()
        self.tree.delete(*self.tree.get_children())
        shown = 0
        for entity in self.entities:
            if any(getattr(entity, field) != value for field, value in selected.items()):
                continue
            if query and query not in _search_blob(entity):
                continue
            tag = "even" if shown % 2 == 0 else "odd"
            self.tree.insert("", tk.END, values=_entity_values(entity), tags=(tag,))
            shown += 1
        self.count_label.config(text=f"Showing {shown} of {len(self.entities)}")


if __name__ == "__main__":
    root = tk.Tk()
    apply_theme(root, "light")
    root.title("Learning Entities")
    root.geometry("1180x640")
    page = EntityPage(root)
    page.pack(fill=tk.BOTH, expand=True)
    theme_btn = ttk.Button(root, text=toggle_label(), style="Back.TButton")

    def on_toggle():
        toggle_theme(root)
        theme_btn.config(text=toggle_label())
        page.refresh_theme()
        theme_btn.lift()

    theme_btn.config(command=on_toggle)
    theme_btn.place(relx=1.0, x=-12, y=10, anchor="ne")
    theme_btn.lift()
    root.mainloop()
